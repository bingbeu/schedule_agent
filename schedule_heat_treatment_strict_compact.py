# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import math
import os
from functools import lru_cache
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
try:
    from ortools.constraint_solver import pywrapcp, routing_enums_pb2
except ImportError:
    pywrapcp = routing_enums_pb2 = None


PROJECT = Path(os.environ.get("SCHEDULE_PROJECT_DIR", Path(__file__).resolve().parent))
DEFAULT_INPUT = PROJECT / "data"
DEFAULT_RULE = PROJECT / "rule"
DEFAULT_TOOLING = PROJECT / "工模具"
DEFAULT_OUTPUT = PROJECT / "630厂热处理生产计划排产表_结果版_OR双炉排程_严格20格精简版.xlsx"

LIMIT_BLANKS = 20
MIN_PLATFORM_TONNAGE = 50
SMALL_PLATFORM_MAX_ORDERS = 1
HEADER_SCAN_ROWS = 40
STEP_UNIT_MIN = 1 / 60
TIGHT_TEMP = 5
TOOL_TOL = 1e-9  # 只消除浮点误差；现场表未授权0.1mm尺寸容差
THICK_WALL_FALLBACK_MIN = 45
FRONT_STEPS, TEMPER_STEPS = 42, 70


class RuleError(ValueError):
    """缺失或无法解释的工艺规则，不允许自动默认为零空格。"""


def order_id(x):
    return x.get("_uid", norm(x["_contract"]))


def numeric_cell(value, label):
    """严格读取一个数，拒绝多温度、范围及非有限数。"""
    s = text(value).replace("℃", "").strip()
    if not re.fullmatch(r"[+-]?\d+(?:\.\d+)?", s):
        raise ValueError(f"{label}需为单一数值，原值={s!r}")
    n = float(s)
    if not math.isfinite(n):
        raise ValueError(f"{label}必须为有限数")
    return n

ALIASES = {
    "contract": ("合同号", "合同编号", "订单编号", "任务单号", "销售订单"),
    "factory": ("主体厂",),
    "variety": ("品种",),
    "brand": ("牌号",),
    "steel": ("钢级", "钢种"),
    "outer": ("外径", "OD"),
    "wall": ("壁厚", "WT", "厚度"),
    "length": ("长度范围", "长度"),
    "qty": ("计划支数", "支数", "数量"),
    "plan_qty": ("计划产量", "产量"),
    "process": ("热处理方式", "工艺", "工艺路线", "处理方式"),
    "front": ("正火炉温度", "淬火炉温度", "前炉温度", "C1", "淬火温度"),
    "temper": ("回火炉温度", "回火温度", "C2"),
    "speed": ("步进周期", "周期", "节拍"),
    "loading": ("布料方式",),
    "status": ("状态", "生产状态", "备注", "实际生产情况"),
    "urgent": ("急催", "急催备注", "紧急合同标记"),
}

BASE_COLS = [
    "序号", "阶段", "任务ID", "订单编号", "主体厂", "品种", "外径", "长度范围", "热处理方式", "计划产量",
    "牌号", "钢级", "壁厚", "数量", "步进周期", "布料方式", "前炉温度", "回火温度",
    "喷嘴规格", "除鳞环/挡水板规格", "工模具核对状态", "前炉首支进炉", "前炉末支进炉", "前炉末支出炉",
    "回火首支进炉", "回火末支进炉", "回火末支出炉", "状态", "本单开始原因", "是否拥堵",
    "备注", "急催", "急催备注",
]


def norm(x) -> str:
    """规整文本便于匹配。"""
    return str(x).replace("\n", "").replace(" ", "").strip()


def num(x, default=None):
    """提取单元格中的第一个数字。"""
    if pd.isna(x):
        return default
    m = re.search(r"-?\d+(?:\.\d+)?", str(x))
    return float(m.group()) if m else default


def text(x) -> str:
    """返回空值安全文本。"""
    return "" if pd.isna(x) else str(x).strip()


def is_blank(x) -> bool:
    """判断空值或空文本。"""
    return pd.isna(x) or str(x).strip() == ""


def append_note(base, note) -> str:
    """追加备注并去重。"""
    base, note = text(base), text(note)
    if not note or note in base:
        return base
    return f"{base}；{note}" if base else note


def files_from(path: Path) -> list[Path]:
    """展开文件、目录或通配输入。"""
    path = Path(path)
    if path.is_file():
        return [path]
    return [p for p in sorted(path.rglob("*")) if p.suffix.lower() in {".xlsx", ".xls", ".xlsm"} and not p.name.startswith("~$")]


def header_score(vals) -> int:
    """给候选表头行评分。"""
    headers = [norm(v) for v in vals]
    score = 0
    for names in ALIASES.values():
        score += 2 if any(any(norm(a).lower() in h.lower() for a in names) for h in headers) else 0
    return score


def detect_header(df) -> tuple[int, list[str]]:
    """动态识别单行或两行表头。"""
    best = (0, 0, [])
    max_row = min(len(df), HEADER_SCAN_ROWS)
    for i in range(max_row):
        row = [text(v) for v in df.iloc[i].tolist()]
        cand = row
        score = header_score(cand)
        if i + 1 < max_row:
            below = [text(v) for v in df.iloc[i + 1].tolist()]
            merged = [(a + b if a and b and a not in b else a or b) for a, b in zip(row, below)]
            mscore = header_score(merged)
            if mscore > score:
                cand, score = merged, mscore
        if score > best[0]:
            best = (score, i, cand)
    if best[0] < 8:
        raise ValueError("未识别到有效表头")
    headers, seen = [], {}
    for i, h in enumerate(best[2]):
        name = h or f"列{i+1}"
        seen[name] = seen.get(name, 0) + 1
        headers.append(name if seen[name] == 1 else f"{name}.{seen[name]-1}")
    return best[1], headers


def col_by_alias(headers, key):
    """按别名找到列名。"""
    for h in headers:
        if norm(h).lower() in {norm(a).lower() for a in ALIASES[key]}:
            return h
    for h in headers:
        hn = norm(h).lower()
        if any(norm(a).lower() in hn for a in ALIASES[key]):
            return h
    return None


def read_inputs(path: Path) -> tuple[list[dict], list[str]]:
    """读取多个排产 Excel，并保留原始字段。"""
    rows, source_headers = [], []
    for file in files_from(path):
        with pd.ExcelFile(file) as xls:
            for sheet in xls.sheet_names:
                raw = pd.read_excel(xls, sheet_name=sheet, header=None, dtype=object)
                try:
                    hrow, headers = detect_header(raw)
                except Exception:
                    continue
                df = raw.iloc[hrow + 1:].copy()
                df.columns = headers[:len(df.columns)]
                source_headers.extend([h for h in headers if h not in source_headers])
                cmap = {k: col_by_alias(df.columns, k) for k in ALIASES}
                for ridx, r in df.iterrows():
                    contract = text(r.get(cmap["contract"])) if cmap["contract"] else ""
                    proc = text(r.get(cmap["process"])) if cmap["process"] else ""
                    if not contract and not proc:
                        continue
                    order = {h: r.get(h, "") for h in df.columns}
                    for k, c in cmap.items():
                        order[f"_{k}"] = r.get(c) if c else ""
                    order["_contract"] = contract
                    order["_process"] = proc
                    order["_front"] = num(order["_front"])
                    order["_temper"] = num(order["_temper"])
                    order["_outer"] = num(order["_outer"])
                    order["_wall"] = num(order["_wall"])
                    order["_qty"] = int(round(num(order["_qty"], None) or num(order.get("_plan_qty"), 0) or 0))
                    order["_plan_tons"] = num(order.get("_plan_qty"), 0) or 0
                    order["_speed"] = num(order["_speed"])
                    order["_needs_temper"] = "回火" in proc
                    order["_needs_front"] = any(k in proc for k in ("正火", "淬火", "调质")) and not (proc.strip() == "回火返工")
                    order["_loading"] = "间隔布料" if "间隔" in text(order["_loading"]) else "连续布料"
                    order["_note"] = "" if cmap["loading"] else "未提供布料方式，默认连续"
                    order["_source_file"] = file.name
                    order["_source_sheet"] = sheet
                    order["_source_row"] = ridx + 1
                    order["_uid"] = f"{file.name}::{sheet}::{ridx+1}"
                    errors = []
                    for field in ("wall", "speed", "qty"):
                        try:
                            value = numeric_cell(r.get(cmap[field]), field)
                            if value <= 0 or (field == "qty" and not value.is_integer()):
                                raise ValueError(f"{field}必须为正数，支数必须为整数")
                            order["_"+field] = int(value) if field == "qty" else value
                        except ValueError as exc:
                            order["_"+field] = None
                            errors.append(str(exc))
                    if proc == "调质":
                        order["_needs_temper"] = True
                    if not (order["_needs_front"] or order["_needs_temper"]):
                        errors.append(f"未识别热处理路线:{proc}")
                    for field, required in (("front", order["_needs_front"]), ("temper", order["_needs_temper"])):
                        if required:
                            try:
                                value = numeric_cell(r.get(cmap[field]), field)
                                if value <= 0:
                                    raise ValueError(f"{field}温度必须为正数")
                                order["_"+field] = value
                            except ValueError as exc:
                                order["_"+field] = None
                                errors.append(str(exc))
                    if cmap["plan_qty"] and not is_blank(r.get(cmap["plan_qty"])):
                        try:
                            order["_plan_tons"] = numeric_cell(r.get(cmap["plan_qty"]), "计划产量")
                            if order["_plan_tons"] < 0:
                                errors.append("计划产量不能为负数")
                        except ValueError as exc:
                            order["_plan_tons"] = 0
                            errors.append(str(exc))
                    if not contract:
                        errors.append("缺少订单编号")
                    order["_input_errors"] = errors
                    missing = []
                    for label, key in [("订单编号", "_contract"), ("热处理方式", "_process"), ("计划支数", "_qty"), ("步进周期", "_speed"), ("壁厚", "_wall")]:
                        if is_blank(order.get(key)) or order.get(key) == 0:
                            missing.append(label)
                    if order["_needs_front"] and order["_front"] is None:
                        missing.append("前炉温度")
                    if order["_needs_temper"] and order["_temper"] is None:
                        missing.append("回火温度")
                    if missing and not order["_input_errors"]:
                        order["_input_errors"].append(f"缺少:{','.join(missing)}")
                    rows.append(order)
    if not rows:
        raise ValueError("没有读取到可排订单")
    return rows, source_headers


def special_tokens(rule_dir: Path) -> set[str]:
    """读取特殊钢级清单。"""
    out = set()
    for file in files_from(rule_dir):
        with pd.ExcelFile(file) as xls:
            for sheet in xls.sheet_names:
                df = pd.read_excel(xls, sheet_name=sheet, header=None, dtype=object)
                if "特殊钢级清单" not in sheet:
                    continue
                for v in df.iloc[1:, 0]:
                    t = norm(text(v)).upper()
                    if re.fullmatch(r"[A-Z0-9]+(?:-[A-Z0-9]+)*", t):
                        out.add(t)
    if not out:
        raise RuleError("未找到独立的特殊钢级清单，请使用现场确认版规则表")
    return out


def read_rules(rule_dir: Path) -> tuple[list[dict], int, int, set[str]]:
    """读取换规规则和炉台步数。"""
    rules, front_steps, temper_steps = [], None, None
    specials = special_tokens(rule_dir)
    for file in files_from(rule_dir):
        with pd.ExcelFile(file) as xls:
            for sheet in xls.sheet_names:
                df = pd.read_excel(xls, sheet_name=sheet, header=None, dtype=object)
                flat = " ".join(text(v) for v in df.to_numpy().ravel())
                for _, row in df.iterrows():
                    vals = [text(v) for v in row.tolist()]
                    if "总料位" in vals:
                        j = vals.index("总料位")
                        f, t = numeric_cell(row.iloc[j+1], "前炉料位"), numeric_cell(row.iloc[j+2], "回火料位")
                        if min(f,t) <= 0 or not (f.is_integer() and t.is_integer()):
                            raise RuleError("炉台料位必须为正整数")
                        if front_steps is not None and (front_steps,temper_steps) != (int(f),int(t)):
                            raise RuleError("规则文件中的炉台料位不一致")
                        front_steps, temper_steps = int(f), int(t)
                for i, row in df.iterrows():
                    vals = [text(v) for v in row.tolist()]
                    if not ("条件1" in vals and any("回火" in v and "空格" in v for v in vals)):
                        continue
                    cols = {v: j for j, v in enumerate(vals)}
                    c1 = vals.index("条件1")
                    c2 = vals.index("条件2") if "条件2" in vals else c1 + 1
                    tb = next(j for j, v in enumerate(vals) if "回火" in v and "空格" in v)
                    fb = next(j for j, v in enumerate(vals) if ("正火" in v or "淬火" in v or "前炉" in v) and "空格" in v)
                    for _, r in df.iloc[i + 1:].iterrows():
                        cond = text(r.iloc[c1])
                        mat = text(r.iloc[c2])
                        if not cond or "说明" in cond:
                            continue
                        fblank, tblank = num(r.iloc[fb]), num(r.iloc[tb])
                        if fblank is None and tblank is None:
                            continue
                        rules.append({
                            "序号": len(rules) + 1,
                            "条件1": cond,
                            "条件2": mat,
                            "前炉空格": fblank,
                            "回火空格": tblank,
                            "来源": f"{file.name}/{sheet}/行{int(r.name)+1}",
                        })
    if not rules:
        raise ValueError("未读取到换规规则")
    if front_steps is None or temper_steps is None:
        raise ValueError("未读取到前炉/回火炉步数")
    return rules, front_steps, temper_steps, specials


def atom_ok(atom: str, dt: float, dw: float) -> bool:
    """判断一个 ΔC/ΔT 条件是否成立。"""
    s = atom.replace(" ", "").replace("＜", "<").replace("≤", "<=").replace("≥", ">=").replace("ΔC", "C").replace("ΔT", "T")
    chain = re.fullmatch(r"(-?\d+(?:\.\d+)?)(<=|<)(C|T)(<=|<)(-?\d+(?:\.\d+)?)", s)
    if chain:
        lo, lop, var, rop, hi = chain.groups()
        val = dt if var == "C" else dw
        left = float(lo) <= val if lop == "<=" else float(lo) < val
        right = val <= float(hi) if rop == "<=" else val < float(hi)
        return left and right
    m = re.fullmatch(r"(C|T)(<=|>=|<|>|=)(-?\d+(?:\.\d+)?)", s)
    if not m:
        m = re.fullmatch(r"(-?\d+(?:\.\d+)?)(<=|>=|<|>|=)(C|T)", s)
        if not m:
            return False
        n, op, var = float(m.group(1)), m.group(2), m.group(3)
        val = dt if var == "C" else dw
        return {"<": n < val, "<=": n <= val, ">": n > val, ">=": n >= val, "=": abs(n - val) < 1e-9}[op]
    var, op, n = m.group(1), m.group(2), float(m.group(3))
    val = dt if var == "C" else dw
    return {"<": val < n, "<=": val <= n, ">": val > n, ">=": val >= n, "=": abs(val - n) < 1e-9}[op]


def condition_ok(cond: str, temp_diff: float, wall_diff: float) -> bool:
    """只有显式的或/OR表示 OR；逗号、分号和且均表示 AND。"""
    cond = cond.replace("并且", "且")
    for alt in re.split(r"或|\bor\b", cond, flags=re.I):
        atoms = [a for a in re.split(r"且|and|[,，;；]", alt, flags=re.I) if a.strip()]
        if atoms and all(atom_ok(a, temp_diff, wall_diff) for a in atoms):
            return True
    return False


def material_ok(rule_mat: str, nxt: dict, specials: set[str]) -> bool:
    """按后订单钢级判断条件2。"""
    mat = text(rule_mat)
    rear = f"{text(nxt.get('_steel'))} {text(nxt.get('_brand'))}".upper()
    tokens = set(re.findall(r"[A-Z0-9]+(?:-[A-Z0-9]+)*", rear))
    is_special = bool(tokens & specials)
    if "其它" in mat:
        return not is_special
    if "特殊" in mat:
        return is_special
    return not mat or mat.upper() in rear


def blank_count(prev: dict | None, cur: dict, furnace: str, rules, specials) -> tuple[float, str]:
    """计算相邻订单指定炉台换规空格。"""
    if prev is None:
        return 0, "首单"
    if furnace == "前炉":
        if not (prev["_needs_front"] and cur["_needs_front"]):
            return 0, "无前炉相邻"
        temp_diff = abs(prev["_front"] - cur["_front"])
        key = "前炉空格"
    else:
        if not (prev["_needs_temper"] and cur["_needs_temper"]):
            return 0, "无回火相邻"
        temp_diff = abs(prev["_temper"] - cur["_temper"])
        key = "回火空格"
    wall_diff = abs(prev["_wall"] - cur["_wall"])
    rule_rows=tuple((r["序号"],r["条件1"],r["条件2"],r[key]) for r in rules)
    rear=f"{text(cur.get('_steel'))} {text(cur.get('_brand'))}".upper()
    value,number=_blank_lookup(rule_rows,tuple(sorted(specials)),rear,temp_diff,wall_diff,furnace)
    return value, f"第{number}条，ΔC={temp_diff:g}℃，ΔT={wall_diff:g}"


@lru_cache(maxsize=20000)
def _blank_lookup(rule_rows, specials, rear, temp_diff, wall_diff, furnace):
    hits=[(seq,value) for seq,cond,mat,value in rule_rows if material_ok(mat,{"_steel":rear,"_brand":""},set(specials)) and condition_ok(cond,temp_diff,wall_diff)]
    if not hits:
        raise RuleError(f"{furnace}未命中规则，ΔC={temp_diff:g}℃，ΔT={wall_diff:g}")
    if any(value is None for _,value in hits):
        raise RuleError(f"{furnace}命中规则的空格数缺失，需现场确认")
    number,value=max(hits,key=lambda r:r[1])
    return value,number


def blank_text(value: float) -> str:
    """把空格数转成车间易读文本。"""
    return f"{value:g}格" + ("，超20" if value > LIMIT_BLANKS else "")


def edge_blank_summary(prev: dict | None, cur: dict | None, rules, specials) -> tuple[float, str]:
    """汇总相邻订单双炉最大空格和简短说明。"""
    if prev is None or cur is None:
        return 0, "无相邻订单"
    parts, values = [], []
    if prev["_needs_front"] and cur["_needs_front"]:
        fb, fr = blank_count(prev, cur, "前炉", rules, specials)
        parts.append(f"前炉空{fb:g}格（{fr}）")
        values.append(fb)
    if prev["_needs_temper"] and cur["_needs_temper"]:
        tb, tr = blank_count(prev, cur, "回火炉", rules, specials)
        parts.append(f"回火空{tb:g}格（{tr}）")
        values.append(tb)
    return (max(values) if values else 0), "，".join(parts) if parts else "无同炉相邻"


def edge_ok(state, cur, rules, specials) -> tuple[bool, str]:
    """判断订单接到当前主序列后是否超过20空格或温度回摆。"""
    pf, pt = state.get("front"), state.get("temper")
    if cur["_needs_front"] and not cur["_needs_temper"] and pt is not None and state.get("front_inserted", 0) >= 1:
        return False, "两个回火订单之间已插入1单仅前炉订单，继续插入会造成回火炉空炉风险"
    try:
        fb, fr = blank_count(pf, cur, "前炉", rules, specials) if cur["_needs_front"] else (0, "无")
        tb, tr = blank_count(pt, cur, "回火炉", rules, specials) if cur["_needs_temper"] else (0, "无")
    except RuleError as exc:
        return False, str(exc)
    if pf and cur["_needs_front"] and abs(abs(cur["_front"]-pf["_front"])-60) < 1e-9:
        material = set(re.findall(r"[A-Z0-9]+(?:-[A-Z0-9]+)*", f"{text(cur.get('_steel'))} {text(cur.get('_brand'))}".upper()))
        if material & {"4140", "35CRMO", "42CRMO", "45", "45钢"}:
            return False, "现场规则备注6:该钢级淬火温差60℃需空炉换规，空格折算尚未确认"
    front_cross = bool(pf and cur["_needs_front"] and abs(cur["_front"] - pf["_front"]) > TIGHT_TEMP)
    temper_rise_cross = bool(pt and cur["_needs_temper"] and cur["_temper"] > pt["_temper"] + TIGHT_TEMP)
    if fb > LIMIT_BLANKS and not front_cross:
        return False, f"前炉空{fb:g}格，超过{LIMIT_BLANKS}格"
    if tb > LIMIT_BLANKS and not temper_rise_cross:
        return False, f"回火空{tb:g}格，超过{LIMIT_BLANKS}格"
    if cur["_needs_temper"] and pt and cur["_temper"] + TIGHT_TEMP < pt["_temper"]:
        return False, f"回火温度回摆：{pt['_temper']:g}->{cur['_temper']:g}"
    if cur["_needs_front"] and not cur["_needs_temper"] and pf and state.get("front_prev"):
        a, b, c = state["front_prev"]["_front"], pf["_front"], cur["_front"]
        if abs(b - a) > TIGHT_TEMP and abs(c - b) > TIGHT_TEMP and (b - a) * (c - b) < 0:
            return False, f"前炉温度回摆：{a:g}->{b:g}->{c:g}"
    return True, "连续接料，无额外空炉"


def seed_order_with_ortools(items: list[dict]) -> list[dict]:
    """用 OR-Tools 生成回火升温初始序列。"""
    if len(items) <= 2 or pywrapcp is None:
        return sorted(items, key=lambda x: (x["_temper"] if x["_needs_temper"] else 9999, x["_front"] or 9999, x["_wall"]))
    manager = pywrapcp.RoutingIndexManager(len(items), 1, 0)
    routing = pywrapcp.RoutingModel(manager)

    def cost(i, j):
        a, b = items[manager.IndexToNode(i)], items[manager.IndexToNode(j)]
        at = a["_temper"] if a["_needs_temper"] else a.get("_front", 0)
        bt = b["_temper"] if b["_needs_temper"] else b.get("_front", 0)
        down = max(0, at - bt - TIGHT_TEMP)
        return int((bt - at) ** 2 + 50000 * down ** 2 + 20 * abs((a.get("_front") or 0) - (b.get("_front") or 0)) ** 2 + 100 * abs(a["_wall"] - b["_wall"]))

    idx = routing.RegisterTransitCallback(cost)
    routing.SetArcCostEvaluatorOfAllVehicles(idx)
    params = pywrapcp.DefaultRoutingSearchParameters()
    params.first_solution_strategy = routing_enums_pb2.FirstSolutionStrategy.PATH_CHEAPEST_ARC
    params.local_search_metaheuristic = routing_enums_pb2.LocalSearchMetaheuristic.GUIDED_LOCAL_SEARCH
    params.time_limit.seconds = 8
    sol = routing.SolveWithParameters(params)
    if sol is None:
        return sorted(items, key=lambda x: (x["_temper"] if x["_needs_temper"] else 9999, x["_front"] or 9999, x["_wall"]))
    out, cur = [], routing.Start(0)
    while not routing.IsEnd(cur):
        out.append(items[manager.IndexToNode(cur)])
        cur = sol.Value(routing.NextVar(cur))
    return out


def build_sequence(items, rules, specials) -> tuple[list[dict], list[dict]]:
    """构造严格20格主序列，不能接续的下沉。"""
    pool = list(items)
    pool.sort(key=lambda x: (x["_temper"] if x["_needs_temper"] else 10_000, x["_front"] or 10_000, x["_wall"], x["_speed"]))
    main, deferred, state = [], [], {"front": None, "front_prev": None, "temper": None, "front_inserted": 0}
    while pool:
        candidates = []
        for x in pool:
            ok, reason = edge_ok(state, x, rules, specials)
            if not ok:
                continue
            fb = blank_count(state.get("front"), x, "前炉", rules, specials)[0] if x["_needs_front"] else 0
            tb = blank_count(state.get("temper"), x, "回火炉", rules, specials)[0] if x["_needs_temper"] else 0
            t = x["_temper"] if x["_needs_temper"] else (state["temper"]["_temper"] if state.get("temper") else 9999)
            front_only = 1 if (x["_needs_front"] and not x["_needs_temper"]) else 0
            candidates.append((front_only, t, tb, fb, abs((x.get("_front") or 0) - (state.get("front") or {}).get("_front", x.get("_front") or 0)), x["_wall"], x, reason))
        if not candidates:
            for x in pool:
                ok, reason = edge_ok(state, x, rules, specials)
                x["_defer_reason"] = f"不能接入当前集中生产路径：{reason}"
                deferred.append(x)
            break
        *_, chosen, reason = min(candidates, key=lambda v: v[:6])
        chosen["_start_reason"] = reason
        main.append(chosen)
        pool.remove(chosen)
        if chosen["_needs_front"]:
            state["front_prev"], state["front"] = state.get("front"), chosen
        if chosen["_needs_temper"]:
            state["temper"] = chosen
            state["front_inserted"] = 0
        elif chosen["_needs_front"]:
            state["front_inserted"] = state.get("front_inserted", 0) + 1
    return main, deferred


def platform_key(x):
    """返回双炉同温区键。"""
    if not (x["_needs_front"] and x["_needs_temper"]):
        return None
    return (x["_front"], x["_temper"])


def insert_by_temperature(main, group):
    """把同温区订单插入回火升温位置。"""
    t = group[0]["_temper"]
    pos = next((i for i, x in enumerate(main) if x["_needs_temper"] and x["_temper"] > t + TIGHT_TEMP), len(main))
    same = [i for i, x in enumerate(main) if platform_key(x) == platform_key(group[0])]
    if same:
        pos = max(same) + 1
    return main[:pos] + group + main[pos:]


def restore_large_platforms(main, deferred, rules=None, specials=None):
    """把同温区累计产量达到阈值的订单恢复到正常生产段。"""
    if rules is None or specials is None:
        raise ValueError("大批量恢复必须传入实际规则，并校验候选")
    from engine import sequence_problems
    groups = {}
    for x in deferred:
        key = platform_key(x)
        if key:
            groups.setdefault(key, []).append(x)
    restored = set()
    for key, arr in groups.items():
        tons = sum(float(x.get("_plan_tons", 0) or 0) for x in arr)
        if tons < MIN_PLATFORM_TONNAGE:
            continue
        arr.sort(key=lambda x: (x["_wall"], x["_speed"], x["_contract"]))
        candidate = insert_by_temperature(main, arr)
        if sequence_problems(candidate, rules, specials):
            continue
        for x in arr:
            x["_platform_tons"] = tons
            x["_start_reason"] = f"同温区累计计划产量{tons:g}吨≥{MIN_PLATFORM_TONNAGE}吨，恢复到正常集中生产"
            restored.add(id(x))
        main = candidate
    return main, [x for x in deferred if id(x) not in restored]


def insert_front_order_near_temper_only(main, deferred, rules, specials):
    """在只回火订单前插入一单不会超过20空格的仅前炉订单。"""
    moved = set()
    i = 0
    while i < len(main):
        cur = main[i]
        if not (cur["_needs_temper"] and not cur["_needs_front"]):
            i += 1
            continue
        prev_front = next((x for x in reversed(main[:i]) if x["_needs_front"]), None)
        next_front = next((x for x in main[i + 1:] if x["_needs_front"]), None)
        candidates = []
        for x in deferred:
            if id(x) in moved or not (x["_needs_front"] and not x["_needs_temper"]):
                continue
            left_blank, left_rule = blank_count(prev_front, x, "前炉", rules, specials)
            right_blank, right_rule = blank_count(x, next_front, "前炉", rules, specials) if next_front else (0, "尾段")
            if left_blank > LIMIT_BLANKS or right_blank > LIMIT_BLANKS:
                continue
            score = (
                abs((prev_front["_front"] if prev_front else x["_front"]) - x["_front"])
                + abs((next_front["_front"] if next_front else x["_front"]) - x["_front"]),
                left_blank + right_blank,
                abs(x["_wall"] - (prev_front["_wall"] if prev_front else x["_wall"])),
            )
            candidates.append((score, x, left_blank, left_rule, right_blank, right_rule))
        if candidates:
            _, x, lb, lr, rb, rr = min(candidates, key=lambda v: v[0])
            x["_start_reason"] = f"只回火期间补前炉：前侧{lb:g}格，后侧{rb:g}格，均≤{LIMIT_BLANKS}"
            main.insert(i, x)
            moved.add(id(x))
            i += 2
        else:
            msg = f"只回火：未找到温度相近且两侧空格≤{LIMIT_BLANKS}的前炉订单搭配"
            cur["_start_reason"] = msg
            cur["_note"] = append_note(cur.get("_note", ""), msg)
            i += 1
    return main, [x for x in deferred if id(x) not in moved]


def defer_small_islands(main, deferred, rules, specials):
    """下沉前后都需大空格的小批孤岛温区。"""
    changed = True
    while changed:
        changed, groups = False, {}
        for i, x in enumerate(main):
            key = platform_key(x)
            if key:
                groups.setdefault(key, []).append(i)
        remove = {}
        for key, idxs in groups.items():
            if len(idxs) > SMALL_PLATFORM_MAX_ORDERS:
                continue
            tons = sum(float(main[i].get("_plan_tons", 0) or 0) for i in idxs)
            if tons >= MIN_PLATFORM_TONNAGE:
                continue
            first, last = idxs[0], idxs[-1]
            left_blank, left_detail = edge_blank_summary(main[first - 1] if first else None, main[first], rules, specials)
            right_blank, right_detail = edge_blank_summary(main[last], main[last + 1] if last + 1 < len(main) else None, rules, specials)
            if left_blank > LIMIT_BLANKS and right_blank > LIMIT_BLANKS:
                for i in idxs:
                    main[i]["_defer_reason"] = f"暂缓：小批温区，{len(idxs)}单/{tons:g}吨，前后空格都超过{LIMIT_BLANKS}"
                    main[i]["_note"] = append_note(
                        main[i].get("_note", ""),
                        f"小批温区暂缓：前侧{left_detail}；后侧{right_detail}",
                    )
                    remove[i] = main[i]
        if remove:
            deferred.extend(remove[i] for i in sorted(remove))
            main = [x for i, x in enumerate(main) if i not in remove]
            changed = True
    return main, deferred


def annotate_reasons(main, rules, specials):
    """给正常生产订单写清连续性和空格数。"""
    pf = pt = None
    for x in main:
        front_blank = temper_blank = None
        front_label = temper_label = None
        if x["_needs_front"]:
            fb, fr = blank_count(pf, x, "前炉", rules, specials)
            if pf is None:
                front_label = "前炉首单"
            else:
                front_blank = fb
                front_label = f"前炉{blank_text(fb)}"
            pf = x
        if x["_needs_temper"]:
            tb, tr = blank_count(pt, x, "回火炉", rules, specials)
            if pt is None:
                temper_label = "回火首单"
            else:
                temper_blank = tb
                temper_label = f"回火{blank_text(tb)}"
            pt = x
        extra = x.get("_start_reason", "")
        active = [v for v in (front_blank, temper_blank) if v is not None]
        labels = "，".join(v for v in (front_label, temper_label) if v)
        if x.get("_platform_tons"):
            reason = f"同温区大批量：{x['_platform_tons']:g}吨，正常排产；{labels}"
        elif "补前炉" in extra:
            reason = f"{extra}；{labels}"
        elif "只回火" in extra:
            reason = f"{extra}；{labels}"
        elif not active:
            reason = f"首单启动：{labels}"
        elif max(active) <= LIMIT_BLANKS:
            reason = f"正常连续：{labels}"
        else:
            reason = f"必要换规：{labels}"
        x["_start_reason"] = reason
    return main


def clock(minutes) -> str:
    """分钟转第几天 HH:MM:SS。"""
    if minutes is None or minutes == "":
        return ""
    sec = int(round(float(minutes) * 60))
    day, rem = divmod(sec, 24 * 3600)
    h, rem = divmod(rem, 3600)
    m, s = divmod(rem, 60)
    prefix = "" if day == 0 else f"第{day+1}天 "
    return f"{prefix}{h:02d}:{m:02d}:{s:02d}"


def schedule_times(main, rules, specials, front_steps, temper_steps) -> list[dict]:
    """计算双炉首末支进出炉时间。"""
    front_ready = temper_ready = 0.0
    pf = pt = None
    for i, x in enumerate(main, 1):
        step = x["_speed"] * STEP_UNIT_MIN
        gap = step * (2 if x["_loading"] == "间隔布料" else 1)
        qty = x["_qty"]
        front_hold, temper_hold = front_steps * step, temper_steps * step
        fs = ts = None
        if x["_needs_front"]:
            fb, fr = blank_count(pf, x, "前炉", rules, specials)
            anchor = 0 if pf is None else pf["_front_start"] + (pf["_qty"] - 1) * pf["_gap"]
            fs = max(front_ready, anchor + fb * step)
        if x["_needs_temper"]:
            tb, tr = blank_count(pt, x, "回火炉", rules, specials)
            anchor = 0 if pt is None else pt["_temper_start"] + (pt["_qty"] - 1) * pt["_gap"]
            ts = max(temper_ready, anchor + tb * step)
            if x["_needs_front"]:
                ts = max(ts, fs + front_hold)
        if x["_needs_front"] and not x["_needs_temper"]:
            ts = None
        if not x["_needs_front"] and x["_needs_temper"]:
            fs = None
        x.update({"_seq": i, "_gap": gap, "_front_start": fs, "_temper_start": ts, "_deferred": False})
        # 在保持回火最早可行时间的前提下，将前炉首支推迟到及时供料。
        if fs is not None and ts is not None:
            fs = max(fs, ts - front_hold)
            x["_front_start"] = fs
        x["_front_hold"], x["_temper_hold"] = front_hold, temper_hold
        x.pop("_defer_reason", None)
        if fs is not None:
            x["_front_last_in"] = fs + (qty - 1) * gap
            x["_front_end"] = x["_front_last_in"] + front_hold
            front_ready = fs + qty * gap
            pf = x
        if ts is not None:
            x["_temper_last_in"] = ts + (qty - 1) * gap
            x["_temper_end"] = x["_temper_last_in"] + temper_hold
            temper_ready = ts + qty * gap
            pt = x
    return main


def parse_range(v):
    """解析 Φa~Φb 数值范围。"""
    nums = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", text(v))]
    if not nums:
        return None
    return (min(nums), max(nums)) if len(nums) > 1 else (nums[0], nums[0])


def in_range(v, span, tol=TOOL_TOL):
    """带容差范围判断。"""
    return v is not None and span and span[0] - tol <= float(v) <= span[1] + tol


def tooling_formula(text_block: str):
    """从工模具说明中读取条件2备用内径公式。"""
    s = norm(text_block).replace("×", "*").replace("＊", "*")
    m = re.search(r"备用内径=.*?外径-(\d+(?:\.\d+)?)\*壁厚-(\d+(?:\.\d+)?)", s)
    return (float(m.group(1)), float(m.group(2))) if m else (None, None)


def read_tooling(path: Path) -> tuple[list[dict], list[dict]]:
    """读取真实规格及来源；空库存不解释为0或不可用。"""
    nozzles, descales = [], []
    headers = ("管坯外径(mm)", "管坯内径(mm)", "喷嘴内径d(mm)", "分水锥直径e(mm)")
    for file in files_from(path):
        with pd.ExcelFile(file) as xls:
            for sheet in xls.sheet_names:
                df = pd.read_excel(xls, sheet_name=sheet, header=None, dtype=object)
                full = "\n".join(text(v) for v in df.to_numpy().ravel())
                coef, offset = tooling_formula(full)
                for i, row in df.iterrows():
                    vals = [norm(text(v)).replace("（", "(").replace("）", ")") for v in row.tolist()]
                    if all(h in vals for h in headers):
                        od, inn, nz, cone = (vals.index(h) for h in headers)
                        stock = vals.index("数量") if "数量" in vals else None
                        for j, r in df.iloc[i + 1:].iterrows():
                            if not text(r.iloc[od]):
                                continue
                            outer, inner = parse_range(r.iloc[od]), parse_range(r.iloc[inn])
                            nozzle, cone_text = text(r.iloc[nz]), text(r.iloc[cone])
                            source = f"{file.name}::{sheet}::{j + 1}"
                            if not outer or not inner or not nozzle or not cone_text:
                                raise RuleError(f"工模具规则字段不完整:{source}")
                            nozzles.append({
                                "外径": outer, "内径": inner, "喷嘴": nozzle, "分水锥": cone_text,
                                "条件2系数": coef, "条件2扣减": offset,
                                "来源": source, "库存原值": text(r.iloc[stock]) if stock is not None else "",
                            })
                for m in re.finditer(r"(\d+)\.\s*(\d+(?:\.\d+)?)\s*([≤＜<])\s*[φΦ]\s*([≤＜<])\s*(\d+(?:\.\d+)?)", full):
                    idx, lo, lop, hop, hi = m.groups()
                    if float(lo) >= float(hi):
                        raise RuleError(f"除鳞环范围错误:{file.name}::{sheet}::{idx}")
                    descales.append({"规格": f"除鳞环{idx}（{lo}{lop}φ{hop}{hi}）", "下": float(lo), "上": float(hi),
                                     "左闭": lop == "≤", "右闭": hop == "≤", "来源": f"{file.name}::{sheet}::条款{idx}"})
    return nozzles, descales


def _tool_description(r):
    return f"{r['喷嘴']}（{r['分水锥']}）"


def _unique_nozzle(rows, explanation):
    """多种不同工具不能用最窄范围或表格顺序擅自裁决。"""
    choices = sorted({_tool_description(r) for r in rows})
    sources = "、".join(r.get("来源", "未标注来源") for r in rows)
    if len(choices) != 1:
        return "未匹配", explanation + "；喷嘴/分水锥选择待确认，候选:" + "、".join(choices) + "；来源:" + sources
    return choices[0], explanation + "；来源:" + sources


def nozzle_match(x, rules) -> tuple[str, str]:
    outer, wall = x.get("_outer"), x.get("_wall")
    if (outer is None or wall is None or not math.isfinite(outer) or not math.isfinite(wall)
            or outer <= 0 or wall <= 0 or outer <= 2 * wall):
        return "未匹配", "外径/壁厚缺失或不满足正内径"
    if not rules:
        return "未匹配", "未提供喷嘴规则"
    inner = outer - 2 * wall
    exact = [r for r in rules if in_range(outer, r["外径"]) and in_range(inner, r["内径"])]
    if exact:
        return _unique_nozzle(exact, f"喷嘴条件1：内径{inner:g}mm")

    formulas = {(r.get("条件2系数"), r.get("条件2扣减")) for r in rules}
    if len(formulas) != 1 or any(v is None for v in next(iter(formulas))):
        return "未匹配", "条件2备用内径公式缺失或不一致"
    coef, offset = next(iter(formulas))
    spare = outer - coef * wall - offset
    if spare <= 0:
        return "未匹配", f"备用内径{spare:g}mm不为正数"
    # 表中外径为范围：暂用外径档下限作为候选值，不额外添加内径过滤。
    # 原说明未指定档内喷嘴选择，多个不同工具一律待确认。
    candidates = [r for r in rules if -TOOL_TOL <= r["外径"][0] - spare <= 8 + TOOL_TOL]
    if not candidates:
        return "未匹配", f"备用内径{spare:g}mm无外径档下限满足向上增量≤8mm（档位解释待现场确认）"
    lower = min(r["外径"][0] for r in candidates)
    selected = [r for r in candidates if r["外径"][0] == lower]
    explanation = f"喷嘴条件2：备用内径{spare:g}mm，候选外径档下限{lower:g}mm，增量{lower-spare:g}mm（档位解释待现场确认）"
    # 外径档的代表值本身尚未获现场确认，即使工具唯一也不能自动放行。
    choice, reason = _unique_nozzle(selected, explanation)
    return "未匹配", reason if choice == "未匹配" else reason + f"；候选:{choice}，条件2需现场确认"


def descale_match(x, rules) -> str:
    """按外径匹配除鳞环。"""
    outer = x["_outer"]
    hits = []
    for r in rules:
        if outer is None:
            continue
        left = outer >= r["下"] if r["左闭"] else outer > r["下"]
        right = outer <= r["上"] if r["右闭"] else outer < r["上"]
        if left and right:
            hits.append(r["规格"])
    return hits[0] if len(set(hits)) == 1 else "未匹配"


def apply_tooling(rows, nozzles, descales):
    """给订单补充工模具规格。"""
    for x in rows:
        nz, note = nozzle_match(x, nozzles)
        x["_nozzle"], x["_descale"] = nz, descale_match(x, descales)
        if x["_descale"] == "未匹配":
            note = append_note(note, "除鳞环未匹配：外径不在唯一有效区间或规则缺失")
        x["_tooling_reason"] = note
        x["_tooling_status"] = "规格待确认" if nz == "未匹配" or x["_descale"] == "未匹配" else "规格匹配，库存待核实"
        if note:
            x["_note"] = append_note(x.get("_note", ""), note)


def overview_row(x, seq=None, deferred=False):
    """生成排产总览行。"""
    if deferred:
        reason = x.get("_defer_reason", "不能接入当前集中生产路径，暂缓/剔除主序列")
        return {
            "序号": seq, "阶段": "暂缓/剔除主序列", "任务ID":order_id(x), "订单编号": x["_contract"], "主体厂": x.get("_factory", ""),
            "品种": x.get("_variety", ""), "外径": x["_outer"], "长度范围": x.get("_length", ""), "热处理方式": x["_process"],
            "计划产量": x.get("_plan_qty", ""), "牌号": x.get("_brand", ""), "钢级": x.get("_steel", ""), "壁厚": x["_wall"],
            "数量": x["_qty"], "步进周期": x["_speed"], "布料方式": x["_loading"], "前炉温度": x.get("_front", ""),
            "回火温度": x.get("_temper", ""), "喷嘴规格": x.get("_nozzle", ""), "除鳞环/挡水板规格": x.get("_descale", ""), "工模具核对状态": x.get("_tooling_status", "规格待确认"),
            "状态": append_note(text(x.get("_status")), "暂缓/剔除主序列"), "本单开始原因": reason,
            "是否拥堵": "未排", "备注": append_note(x.get("_note", ""), reason), "急催": x.get("_urgent", ""), "急催备注": "",
        }
    return {
        "序号": x["_seq"], "阶段": "正常生产", "任务ID":order_id(x), "订单编号": x["_contract"], "主体厂": x.get("_factory", ""),
        "品种": x.get("_variety", ""), "外径": x["_outer"], "长度范围": x.get("_length", ""), "热处理方式": x["_process"],
        "计划产量": x.get("_plan_qty", ""), "牌号": x.get("_brand", ""), "钢级": x.get("_steel", ""), "壁厚": x["_wall"],
        "数量": x["_qty"], "步进周期": x["_speed"], "布料方式": x["_loading"], "前炉温度": x.get("_front", ""),
        "回火温度": x.get("_temper", ""), "喷嘴规格": x.get("_nozzle", ""), "除鳞环/挡水板规格": x.get("_descale", ""), "工模具核对状态": x.get("_tooling_status", "规格待确认"),
        "前炉首支进炉": clock(x.get("_front_start")), "前炉末支进炉": clock(x.get("_front_last_in")), "前炉末支出炉": clock(x.get("_front_end")),
        "回火首支进炉": clock(x.get("_temper_start")), "回火末支进炉": clock(x.get("_temper_last_in")), "回火末支出炉": clock(x.get("_temper_end")),
        "状态": x.get("_status", ""), "本单开始原因": x.get("_start_reason", "连续接料，无额外空炉"),
        "是否拥堵": "否", "备注": x.get("_note", ""), "急催": x.get("_urgent", ""), "急催备注": "",
    }


def pipe_rows(main):
    """生成钢管级时间轴。"""
    rows = []
    for x in main:
        for i in range(x["_qty"]):
            fr = tm = ""
            if x.get("_front_start") is not None:
                a = x["_front_start"] + i * x["_gap"]
                b = a + x.get("_front_hold", x["_speed"] * STEP_UNIT_MIN * FRONT_STEPS)
                fr = f"{clock(a)}-{clock(b)}"
            if x.get("_temper_start") is not None:
                a = x["_temper_start"] + i * x["_gap"]
                b = a + x.get("_temper_hold", x["_speed"] * STEP_UNIT_MIN * TEMPER_STEPS)
                tm = f"{clock(a)}-{clock(b)}"
            rows.append({"排产序号": x["_seq"], "订单编号": x["_contract"], "钢管序号": i + 1, "前炉时段": fr, "回火时段": tm})
    return rows


def write_sheet(wb, name, rows, headers):
    """写入一个简单美观的 sheet。"""
    ws = wb.create_sheet(name)
    ws.append(headers)
    for r in rows:
        if r is None:
            ws.append([""] * len(headers))
        elif isinstance(r, str):
            ws.append([r] + [""] * (len(headers) - 1))
        else:
            ws.append([r.get(h, "") for h in headers])
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    border = Border(bottom=Side(style="thin", color="D9E2F3"))
    for c in ws[1]:
        c.fill, c.font, c.alignment = fill, font, Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows():
        for c in row:
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, h in enumerate(headers, 1):
        width = 12 if h not in {"本单开始原因", "备注"} else 42
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    return ws


def write_excel(path, main, deferred, source_headers):
    """输出排产总览、原始订单数据和钢管级时间轴。"""
    wb = Workbook()
    wb.remove(wb.active)
    overview = [overview_row(x) for x in main]
    if deferred:
        overview += [None, None, "以下订单与此次排程不能集中生产/排产"]
        overview += [overview_row(x, len(main) + i + 1, True) for i, x in enumerate(deferred)]
    write_sheet(wb, "排产总览", overview, BASE_COLS)
    source_cols = ["排产序号", "阶段", "任务ID", "订单编号", "本单开始原因", "状态", "喷嘴规格", "除鳞环/挡水板规格", "工模具核对状态"] + source_headers
    source_rows = []
    for x in main + deferred:
        row = {
            "排产序号": x.get("_seq", ""), "阶段": "正常生产" if not x.get("_defer_reason") else "暂缓/剔除主序列",
            "任务ID":order_id(x), "订单编号": x["_contract"], "本单开始原因": x.get("_defer_reason") or x.get("_start_reason", ""),
            "状态": x.get("_status", ""), "喷嘴规格": x.get("_nozzle", ""), "除鳞环/挡水板规格": x.get("_descale", ""), "工模具核对状态": x.get("_tooling_status", "规格待确认"),
        }
        row.update({h: x.get(h, "") for h in source_headers})
        source_rows.append(row)
    write_sheet(wb, "原始订单数据", source_rows, source_cols)
    write_sheet(wb, "钢管级时间轴", pipe_rows(main), ["排产序号", "订单编号", "钢管序号", "前炉时段", "回火时段"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main():
    """命令行入口。"""
    p = argparse.ArgumentParser()
    p.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    p.add_argument("--change-rules", type=Path, default=DEFAULT_RULE)
    p.add_argument("--tooling-dir", type=Path, default=DEFAULT_TOOLING)
    p.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = p.parse_args()

    from engine import Engine, Schedule
    eng = Engine(args.input, args.change_rules, args.tooling_dir)
    scheduled, deferred = eng.run({})
    Schedule(eng, scheduled, deferred).export(args.output)
    print(f"排产完成：{args.output}")
    print(f"正常生产：{len(scheduled)}单；暂缓/剔除主序列：{len(deferred)}单；订单总数：{len(scheduled)+len(deferred)}")
    print(
        f"前炉步数：{eng.front_steps}；回火炉步数：{eng.temper_steps}；规则空格阈值：{LIMIT_BLANKS}；"
        f"同温区恢复阈值：{MIN_PLATFORM_TONNAGE}吨；小批温区剔除阈值：≤{SMALL_PLATFORM_MAX_ORDERS}单"
    )


if __name__ == "__main__":
    main()

