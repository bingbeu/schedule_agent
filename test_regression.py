"""无需模型密钥的规则、真实输入、动作事务与HTTP回归测试。"""
import copy
import json
import threading
import tempfile
import unittest
import urllib.request
import urllib.error
from pathlib import Path
from unittest.mock import patch

import schedule_heat_treatment_strict_compact as core
from engine import Engine, Schedule, build_sequence, validate, sequence_problems, recover_feasible_insertions
from actions import SchedulerAgent
from agent_cli import dispatch, parse_offline

ROOT = Path(__file__).resolve().parent
APRIL = ROOT / "data" / "4月份排产.xlsx"
JUNE = ROOT / "data" / "630厂热处理车间6月排产.xlsx"


def task(c, temp=500, front=900, wall=10):
    return {"_contract":c,"_uid":c,"_needs_front":True,"_needs_temper":True,
            "_front":front,"_temper":temp,"_wall":wall,"_speed":60,"_qty":3,
            "_loading":"连续布料","_plan_tons":1,"_steel":"P11","_brand":"P11",
            "_outer":100,"_process":"正火+回火","_input_errors":[]}


class SmallEngine(Engine):
    def __init__(self):
        super().__init__(JUNE)
        self.load()
        self.orders=[task("A"),task("B"),task("C")]
        self.source_headers=[]
    def load(self):
        if hasattr(self,"orders"):
            return len(self.orders)
        return super().load()


class RulesAndData(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rules,cls.fs,cls.ts,cls.specials=core.read_rules(ROOT/"rule")

    def test_special_list_exact(self):
        self.assertEqual(self.specials,{"C110","110S","110TS","T95","L80-3CR","L80-1","125","140","4140","SY850"})
        self.assertFalse(core.material_ok("特殊钢级",task("A"),self.specials))

    def test_rule_boundaries_no_gaps_no_overlap(self):
        for dw in (0,6.99,7,13.99,14,14.01,70):
            for dt in (0,9.99,10,10.01,19.99,20,20.01,29.99,30,30.01,49.99,50,50.01,100):
                for steel in ("P11","C110"):
                    x=task("A");x["_steel"]=steel
                    hits=[r for r in self.rules if core.material_ok(r["条件2"],x,self.specials) and core.condition_ok(r["条件1"],dt,dw)]
                    self.assertEqual(len(hits),1,(dw,dt,steel,hits))
        self.assertEqual((self.fs,self.ts),(42,70))

    def test_comma_is_and(self):
        self.assertFalse(core.condition_ok("ΔT<7，ΔC<10",20,3))
        self.assertTrue(core.condition_ok("ΔT<7 或 ΔC<10",20,3))

    def test_rule_missing_rejected(self):
        with self.assertRaises(core.RuleError):core.blank_count(task("A"),task("B"),"前炉",[],self.specials)
        r=copy.deepcopy(self.rules);r[0]["前炉空格"]=None
        with self.assertRaises(core.RuleError):core.blank_count(task("A"),task("B"),"前炉",r,self.specials)

    def test_duplicate_contracts_have_unique_task_ids(self):
        orders,_=core.read_inputs(JUNE)
        self.assertEqual(len(orders),43)
        self.assertEqual(len({core.order_id(x) for x in orders}),43)
        self.assertGreater(len([x for x in orders if x["_contract"]=="2326030045001"]),1)

    def test_multi_temperature_quarantined(self):
        orders,_=core.read_inputs(JUNE)
        bad=[x for x in orders if x["_input_errors"]]
        self.assertEqual(len(bad),1)
        self.assertIsNone(bad[0]["_temper"])

    def test_tonnage_not_converted_to_quantity(self):
        orders,_=core.read_inputs(APRIL)
        bad=[x for x in orders if x["_input_errors"]]
        self.assertEqual(len(bad),8)
        self.assertTrue(all(x["_qty"] is None for x in bad))

    def test_multiple_batches_require_explicit_option(self):
        with self.assertRaises(ValueError):Engine(ROOT/"data").load()
        e=Engine(ROOT/"data",combine_inputs=True);e.load()
        self.assertEqual(len(e.orders),150)

    def test_real_schedules_conserve_tasks_and_pass(self):
        for source,n in ((APRIL,107),(JUNE,43)):
            e=Engine(source);a,b=e.run();s=Schedule(e,a,b)
            self.assertEqual(len(a)+len(b),n)
            self.assertEqual(s.kpi["硬约束违规数"],0)
            self.assertFalse(sequence_problems(a,e.rules,e.specials))
            self.assertTrue(all(not x["_input_errors"] for x in a))
            self.assertEqual({core.order_id(x) for x in a+b},{core.order_id(x) for x in e.orders})

    def test_combined_schedule_conservation(self):
        e=Engine(ROOT/"data",combine_inputs=True);a,b=e.run();s=Schedule(e,a,b)
        self.assertEqual(len(a)+len(b),150)
        self.assertEqual(s.kpi["硬约束违规数"],0)

    def test_recovery_inserts_between_temperatures(self):
        a,b,c=task("A",500),task("B",600),task("C",550)
        main,left,record=recover_feasible_insertions([a,b],[c],self.rules,self.specials)
        self.assertEqual([x["_uid"] for x in main],["A","C","B"])
        self.assertFalse(left);self.assertEqual(record["恢复任务数"],1)

    def test_recovery_preserves_manual_positions(self):
        a,b,c=task("A",500),task("B",600),task("C",550)
        main,left,record=recover_feasible_insertions([a,b],[c],self.rules,self.specials,positions={"B":1})
        self.assertEqual([x["_uid"] for x in main],["A","B"])
        self.assertEqual(len(left),1)

    def test_engine_direct_numeric_validation(self):
        e=Engine(JUNE);e.load()
        with self.assertRaises(ValueError):e.run({"start_delay_minutes":-60})


class Tooling(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.nozzles, cls.descales = core.read_tooling(ROOT / "工模具")

    def test_actual_workbook_and_provenance(self):
        self.assertEqual((len(self.nozzles),len(self.descales)),(27,5))
        self.assertEqual(core.tooling_formula("备用内径 = 外径 - 2 × 壁厚 - 20。"),(2,20))
        self.assertEqual(self.nozzles[0]["来源"],"工模具.xlsx::喷嘴规格::2")
        self.assertEqual(self.nozzles[0]["库存原值"],"")

    def test_condition_one_with_cone(self):
        name, reason = core.nozzle_match({"_outer":559,"_wall":20},self.nozzles)
        self.assertEqual(name,"Φ485（Φ160）")
        self.assertIn("条件1",reason)
        self.assertIn("喷嘴规格::23",reason)

    def test_no_unapproved_dimension_tolerance(self):
        # 原0.1mm容差会错误放行355.6-2*23.83=307.94（上限306/下限308）。
        self.assertEqual(core.nozzle_match({"_outer":355.6,"_wall":23.83},self.nozzles)[0],"未匹配")
        self.assertEqual(core.nozzle_match({"_outer":355.6,"_wall":23.8},self.nozzles)[0],"Φ285（无分水锥）")

    def test_condition_two_not_arbitrary_inner_filter(self):
        # 备用367.34，外径档374增量6.66，但该档5种喷嘴无法唯一确定。
        name,reason=core.nozzle_match({"_outer":406.4,"_wall":9.53},self.nozzles)
        self.assertEqual(name,"未匹配")
        for expected in ("6.66","Φ250","Φ345","选择待确认"):
            self.assertIn(expected,reason)

    def test_condition_two_increment_over_eight(self):
        name,reason=core.nozzle_match({"_outer":406.4,"_wall":10.31},self.nozzles)
        self.assertEqual(name,"未匹配")
        self.assertIn("无外径档下限满足",reason)

    def test_descale_exact_open_closed_boundaries(self):
        for od,index in ((323.8,1),(408,1),(408.0001,2),(457.2,2),(457.2001,3),
                         (510,3),(510.0001,4),(559,4),(559.0001,5),(610,5)):
            self.assertTrue(core.descale_match({"_outer":od},self.descales).startswith(f"除鳞环{index}（"))
        for od in (323,610.0001,None):
            self.assertEqual(core.descale_match({"_outer":od},self.descales),"未匹配")

    def test_overlapping_different_tools_rejected(self):
        rules=copy.deepcopy(self.nozzles[:1]);r=copy.deepcopy(rules[0]);r["喷嘴"]="Φ999";rules.append(r)
        self.assertEqual(core.nozzle_match({"_outer":323,"_wall":53},rules)[0],"未匹配")
        rings=copy.deepcopy(self.descales[:1]);r=copy.deepcopy(rings[0]);r["规格"]="其他除鳞环";rings.append(r)
        self.assertEqual(core.descale_match({"_outer":350},rings),"未匹配")

    def test_real_coverage_is_separate_from_sequence_feasibility(self):
        for source,all_nozzle,all_ring,main_pending in ((APRIL,95,105,12),(JUNE,36,43,5)):
            e=Engine(source);m,d=e.run();s=Schedule(e,m,d)
            self.assertEqual(sum(x["_nozzle"]!="未匹配" for x in m+d),all_nozzle)
            self.assertEqual(sum(x["_descale"]!="未匹配" for x in m+d),all_ring)
            self.assertEqual(s.kpi["主序列工模具待确认单数"],main_pending)
            self.assertEqual(s.kpi["硬约束违规数"],0)
            self.assertTrue(s.tooling_issues)
            self.assertEqual(core.overview_row(m[0])["工模具核对状态"],m[0]["_tooling_status"])


class Actions(unittest.TestCase):
    def setUp(self):
        self.agent=SchedulerAgent(SmallEngine())
        self.agent.reset()

    def test_last_means_last(self):
        self.agent.move("B","last")
        self.assertEqual(core.order_id(self.agent.current.main[-1]),"B")

    def test_first_and_position(self):
        self.agent.move("C","first")
        self.assertEqual(core.order_id(self.agent.current.main[0]),"C")
        self.agent.unpin();self.agent.move("C","2")
        self.assertEqual(core.order_id(self.agent.current.main[1]),"C")

    def test_failed_rerun_preserves_state(self):
        before=copy.deepcopy(self.agent.active);cur=self.agent.current;version=self.agent.version
        with patch.object(self.agent.engine,"run",side_effect=RuntimeError("模拟失败")):
            result=dispatch(self.agent,"temper_change",{"contract":"A","temper":620})
        self.assertIn("error",result)
        self.assertEqual(self.agent.active,before)
        self.assertIs(self.agent.current,cur)
        self.assertEqual(self.agent.version,version)

    def test_infeasible_pin_rejected(self):
        self.agent.engine.orders[0]["_temper"]=600
        before=copy.deepcopy(self.agent.active)
        result=dispatch(self.agent,"pin",{"contracts":["A","B"]})
        self.assertIn("error",result);self.assertEqual(self.agent.active,before)

    def test_no_implicit_extra_furnace(self):
        self.agent.engine.orders[0]["_needs_temper"]=False
        self.assertIn("error",dispatch(self.agent,"temper_change",{"contract":"A","temper":620}))

    def test_bad_numbers_and_unknown_task(self):
        for number in (-1,float("nan"),float("inf"),True):
            self.assertIn("error",dispatch(self.agent,"delay",{"minutes":number}))
        self.assertIn("error",dispatch(self.agent,"temper_change",{"contract":"UNKNOWN","temper":620}))

    def test_allowlist_and_schema(self):
        self.assertIn("error",dispatch(self.agent,"_commit",{}))
        self.assertIn("error",dispatch(self.agent,"boost",None))
        self.assertIn("error",dispatch(self.agent,"remove_orders",{"contracts":"A"}))
        self.assertIn("error",dispatch(self.agent,"export",{"path":"elsewhere"}))

    def test_remove_restore_undo(self):
        self.agent.remove_orders(["A"])
        self.assertEqual(len(self.agent.current.main),2)
        self.agent.restore_orders(["A"])
        self.assertEqual(len(self.agent.current.main),3)
        self.agent.undo();self.assertEqual(len(self.agent.current.main),2)
        self.agent.undo();self.assertEqual(len(self.agent.current.main),3)

    def test_duplicate_contract_rejected(self):
        self.agent.engine.orders[1]["_contract"]="A"
        self.agent.engine.orders[0]["_uid"]="row1"
        self.agent.engine.orders[1]["_uid"]="row2"
        self.assertIn("error",dispatch(self.agent,"explain",{"contract":"A"}))

    def test_list_reflects_changes(self):
        self.agent.temper_change("C",620)
        rows=self.agent.list_orders("C")["订单"]
        self.assertEqual(rows[0]["回火温度"],620)
        self.agent.remove_orders(["C"])
        self.assertEqual(self.agent.list_orders("C")["订单"][0]["阶段"],"已移除")

    def test_export_and_time_validation(self):
        with tempfile.TemporaryDirectory() as directory:
            path=self.agent.current.export(Path(directory)/"result.xlsx")
            from openpyxl import load_workbook
            w=load_workbook(path,data_only=True)
            self.assertIn("钢管级时间轴",w.sheetnames)
            self.assertIn("任务ID",[c.value for c in w["排产总览"][1]])
        self.agent.current.main[0]["_temper_end"]=-1
        self.assertTrue(any(x["类别"]=="违规" for x in validate(self.agent.current.main,self.agent.engine.rules,self.agent.engine.specials,42)))
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(ValueError):self.agent.current.export(Path(directory)/"bad.xlsx")


class Parser(unittest.TestCase):
    def test_negation_and_compound_do_not_mutate(self):
        for text in ("不要移除 A","取消删除 A","延迟 2 小时，然后导出"):
            self.assertIsNone(parse_offline(text)[0][0])

    def test_documented_temperature_command(self):
        self.assertEqual(parse_offline("改 A 回火温度 620")[0],("temper_change",{"contract":"A","temper":620.0}))

    def test_delay_and_last(self):
        self.assertEqual(parse_offline("延迟 2 小时")[0],("delay",{"minutes":120}))
        self.assertEqual(parse_offline("把 A 移到最后")[0],("move",{"contract":"A","position":"last"}))


class Web(unittest.TestCase):
    def setUp(self):
        import webui
        self.web=webui;self.web.engine=SmallEngine();self.web.agent=SchedulerAgent(self.web.engine);self.web.agent.reset()
        self.web.API_KEY="";self.web.HISTORY.clear()

    def test_conversation_history_by_session(self):
        self.web.API_KEY="test"
        calls=[]
        def call(messages,*args):
            calls.append(copy.deepcopy(messages))
            return {"choices":[{"message":{"role":"assistant","content":"测试回复"}}]}
        with patch.object(self.web,"call_llm",side_effect=call):
            self.web.run_turn("解释A","session1")
            self.web.run_turn("把它移到最前","session1")
            self.web.run_turn("KPI","session2")
        self.assertTrue(any(m.get("content")=="解释A" for m in calls[1]))
        self.assertFalse(any(m.get("content")=="解释A" for m in calls[2]))

    def test_http_rejects_bad_json_stale_version_and_cross_origin(self):
        server=self.web.ThreadingHTTPServer(("127.0.0.1",0),self.web.Handler)
        thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        base=f"http://127.0.0.1:{server.server_port}"
        try:
            for body,headers,status in ((b"{bad",{},400),(json.dumps({"text":"KPI","expected_version":-1}).encode(),{},409),(b"{}",{"Origin":"https://example.invalid"},403)):
                req=urllib.request.Request(base+"/api/chat",data=body,headers=headers)
                with self.assertRaises(urllib.error.HTTPError) as error:urllib.request.urlopen(req,timeout=5)
                self.assertEqual(error.exception.code,status)
            req=urllib.request.Request(base+"/api/chat",data=json.dumps({"text":"延迟 2 小时","expected_version":self.web.agent.version}).encode())
            with urllib.request.urlopen(req,timeout=5) as response:
                result=json.load(response)
            self.assertEqual(result["state"]["overrides"]["start_delay_minutes"],120)
        finally:
            server.shutdown();server.server_close();thread.join()


if __name__=="__main__":unittest.main(verbosity=2)
